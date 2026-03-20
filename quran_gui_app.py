import re
import os
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from docx import Document
from docx.shared import RGBColor

# --- Core Logic Functions ---

def hex_to_rgb(hex_code):
    hex_code = hex_code.lstrip('#')
    return RGBColor(int(hex_code[0:2], 16), int(hex_code[2:4], 16), int(hex_code[4:6], 16))

def is_arabic(text):
    return any('\u0600' <= c <= '\u06FF' for c in text)

def process_paragraphs_to_docx(paragraphs, output_docx_path, categories_data):
    """Core engine: takes a list of text lines, applies rules, saves to docx."""
    doc_out = Document()
    
    pattern_parts = []
    keyword_to_color = {}
    
    for cat in categories_data:
        for kw in cat['keywords']:
            keyword_to_color[kw.lower()] = cat['color']
            escaped_kw = re.escape(kw)
            
            if is_arabic(kw):
                pattern_parts.append(escaped_kw) # Arabic: Root word matching
            else:
                pattern_parts.append(r'\b' + escaped_kw + r'\b') # English: Exact word matching
            
    if not pattern_parts:
        return False, "No rules were added to process."

    pattern_str = r'(' + '|'.join(pattern_parts) + r')'
    master_pattern = re.compile(pattern_str, re.IGNORECASE)

    for text in paragraphs:
        text = text.strip()
        if not text:
            doc_out.add_paragraph()
            continue
            
        p = doc_out.add_paragraph()
        
        # If the line is a bracketed tag (e.g. [Arabic] or [Yusuf Ali]), make it bold and skip coloring
        if text.startswith('[') and text.endswith(']'):
            run = p.add_run(text)
            run.font.bold = True
            continue

        last_index = 0
        for match in master_pattern.finditer(text):
            start, end = match.span()
            matched_word = match.group()
            
            if start > last_index:
                p.add_run(text[last_index:start])
                
            run = p.add_run(matched_word)
            run.font.color.rgb = keyword_to_color[matched_word.lower()]
            run.font.bold = True
            last_index = end
            
        if last_index < len(text):
            p.add_run(text[last_index:])

    doc_out.save(output_docx_path)
    return True, ""

def read_file_to_paragraphs(input_path):
    """Helper to extract paragraphs from docx or txt."""
    ext = os.path.splitext(input_path)[1].lower()
    try:
        if ext == '.docx':
            doc_in = Document(input_path)
            return [p.text for p in doc_in.paragraphs], True
        else:
            with open(input_path, 'r', encoding='utf-8') as file:
                return file.readlines(), True
    except Exception as e:
        return [], str(e)

# --- GUI Application Class ---

class QuranColorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Quranic Study Guide - Manuscript Builder")
        self.root.geometry("700x750")
        self.root.configure(padx=20, pady=20)

        self.rules = []
        self.input_files = []
        self.pasted_blocks = [] # Stores tuples of (Label, Text)
        
        # Preloaded Data
        self.PRELOADED_CATEGORIES = [
            "Attributes of Allah", "Legal Commands (Fiqh)", "Prophetic Narratives",
            "Eschatology (The Hereafter)", "Moral & Ethical Imperatives"
        ]
        self.PRELOADED_COLORS = {
            "Deep Blue": "#1A237E", "Forest Green": "#1B5E20", "Goldenrod": "#B8860B",
            "Royal Purple": "#4A148C", "Crimson Red": "#B71C1C", "Saddle Brown": "#5D4037"
        }
        self.TRANSLATION_LABELS = [
            "Arabic", "Yusuf Ali", "Sahih International", "Pickthall", "Commentary"
        ]

        self.build_ui()

    def build_ui(self):
        # --- Section 1: Create Rules ---
        tk.Label(self.root, text="1. Rule Management", font=("Arial", 12, "bold")).pack(anchor="w")
        
        rule_frame = tk.Frame(self.root)
        rule_frame.pack(fill="x", pady=5)
        
        # Dropdowns side-by-side
        tk.Label(rule_frame, text="Category:").grid(row=0, column=0, sticky="w")
        self.cat_var = tk.StringVar()
        ttk.Combobox(rule_frame, textvariable=self.cat_var, values=self.PRELOADED_CATEGORIES, width=25).grid(row=0, column=1, padx=5, sticky="w")

        tk.Label(rule_frame, text="Color:").grid(row=0, column=2, sticky="w")
        self.color_var = tk.StringVar()
        ttk.Combobox(rule_frame, textvariable=self.color_var, values=list(self.PRELOADED_COLORS.keys()), state="readonly", width=15).grid(row=0, column=3, padx=5, sticky="w")

        # Keywords & Add Button
        tk.Label(rule_frame, text="Keywords:").grid(row=1, column=0, sticky="w", pady=5)
        self.kw_entry = tk.Entry(rule_frame, width=45)
        self.kw_entry.grid(row=1, column=1, columnspan=2, padx=5, pady=5, sticky="w")
        tk.Button(rule_frame, text="+ Add Rule", command=self.add_rule).grid(row=1, column=3, padx=5, pady=5)

        self.rules_listbox = tk.Listbox(self.root, height=4)
        self.rules_listbox.pack(fill="x", pady=(0, 15))

        # --- Section 2: Input Management (TABS) ---
        tk.Label(self.root, text="2. Manuscript Inputs", font=("Arial", 12, "bold")).pack(anchor="w")
        
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, pady=5)

        # TAB 1: Multiple File Upload
        self.tab_files = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_files, text="📁 Upload Files")
        
        tk.Button(self.tab_files, text="Browse Multiple Files (.docx, .txt)", command=self.select_files).pack(anchor="w", pady=10, padx=10)
        self.files_listbox = tk.Listbox(self.tab_files)
        self.files_listbox.pack(fill="both", expand=True, padx=10, pady=(0,10))

        # TAB 2: Paste Builder
        self.tab_paste = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_paste, text="📝 Paste & Build Text")
        
        paste_ctrl_frame = tk.Frame(self.tab_paste)
        paste_ctrl_frame.pack(fill="x", padx=10, pady=5)
        
        tk.Label(paste_ctrl_frame, text="Label/Translation:").pack(side="left")
        self.trans_var = tk.StringVar()
        ttk.Combobox(paste_ctrl_frame, textvariable=self.trans_var, values=self.TRANSLATION_LABELS, width=20).pack(side="left", padx=5)
        
        self.paste_area = tk.Text(self.tab_paste, height=6)
        self.paste_area.pack(fill="x", padx=10, pady=5)
        
        tk.Button(self.tab_paste, text="+ Add Block to Manuscript", command=self.add_pasted_block).pack(anchor="e", padx=10)
        
        tk.Label(self.tab_paste, text="Manuscript Preview (Blocks Added):").pack(anchor="w", padx=10, pady=(5,0))
        self.blocks_listbox = tk.Listbox(self.tab_paste, height=4)
        self.blocks_listbox.pack(fill="x", padx=10, pady=(0,10))

        # --- Section 3: Process ---
        tk.Button(self.root, text="GENERATE COLOR CODED DOCUMENT(S)", command=self.run_processor, font=("Arial", 12, "bold"), bg="#4CAF50", fg="black", pady=10).pack(fill="x", pady=10)

    # --- Actions ---

    def add_rule(self):
        cat = self.cat_var.get().strip()
        color_name = self.color_var.get()
        kw_text = self.kw_entry.get().strip()

        if not cat or not color_name or not kw_text:
            messagebox.showwarning("Missing Info", "Please fill out Category, Color, and Keywords.")
            return

        color_hex = self.PRELOADED_COLORS.get(color_name, "#000000")
        keywords = [kw.strip() for kw in kw_text.split(',') if kw.strip()]
        keywords.sort(key=len, reverse=True)

        self.rules.append({'name': cat, 'color': hex_to_rgb(color_hex), 'keywords': keywords})
        self.rules_listbox.insert(tk.END, f"[{cat}] - {color_name} -> {', '.join(keywords)}")
        
        self.kw_entry.delete(0, tk.END)

    def select_files(self):
        filepaths = filedialog.askopenfilenames(
            title="Select Manuscript Drafts",
            filetypes=[("Documents", "*.docx *.txt"), ("All Files", "*.*")]
        )
        for path in filepaths:
            if path not in self.input_files:
                self.input_files.append(path)
                self.files_listbox.insert(tk.END, os.path.basename(path))

    def add_pasted_block(self):
        label = self.trans_var.get().strip()
        text = self.paste_area.get("1.0", tk.END).strip()
        
        if not label or not text:
            messagebox.showwarning("Missing Info", "Please select a label and paste some text.")
            return
            
        self.pasted_blocks.append((label, text))
        preview_text = text[:50] + "..." if len(text) > 50 else text
        self.blocks_listbox.insert(tk.END, f"[{label}] : {preview_text}")
        
        self.paste_area.delete("1.0", tk.END)

    def run_processor(self):
        if not self.rules:
            messagebox.showerror("Error", "Please add at least one rule before processing.")
            return

        active_tab_index = self.notebook.index(self.notebook.select())

        # If User is on the Files Tab
        if active_tab_index == 0:
            if not self.input_files:
                messagebox.showerror("Error", "Please upload at least one file.")
                return
            
            success_count = 0
            for file_path in self.input_files:
                paragraphs, success = read_file_to_paragraphs(file_path)
                if not success:
                    messagebox.showerror("Read Error", f"Failed to read {file_path}:\n{paragraphs}")
                    continue
                
                base_name = os.path.splitext(file_path)[0]
                output_path = f"{base_name}_ColorCoded.docx"
                
                success, msg = process_paragraphs_to_docx(paragraphs, output_path, self.rules)
                if success:
                    success_count += 1
                else:
                    messagebox.showerror("Processing Error", msg)
                    
            messagebox.showinfo("Complete", f"Successfully processed {success_count} file(s).")

        # If User is on the Paste Builder Tab
        elif active_tab_index == 1:
            if not self.pasted_blocks:
                messagebox.showerror("Error", "Please add at least one text block.")
                return
                
            output_path = filedialog.asksaveasfilename(
                defaultextension=".docx",
                initialfile="Compiled_Manuscript_ColorCoded.docx",
                title="Save Compiled Manuscript As",
                filetypes=[("Word Document", "*.docx")]
            )
            
            if not output_path:
                return # User cancelled save dialog

            # Compile the blocks into a list of paragraphs
            compiled_paragraphs = []
            for label, text in self.pasted_blocks:
                compiled_paragraphs.append(f"[{label}]") # Add the bracketed tag automatically
                compiled_paragraphs.extend(text.split('\n')) # Add the pasted text
                compiled_paragraphs.append("") # Add a blank line for spacing

            success, msg = process_paragraphs_to_docx(compiled_paragraphs, output_path, self.rules)
            
            if success:
                messagebox.showinfo("Success!", f"Compiled document saved to:\n{output_path}")
            else:
                messagebox.showerror("Error", msg)

# --- Launch Application ---
if __name__ == "__main__":
    root = tk.Tk()
    app = QuranColorApp(root)
    root.mainloop()
